from io import BytesIO
from pathlib import Path

from docx import Document
from pypdf import PdfReader


MAX_FILE_SIZE_MB = 5
MAX_FILE_SIZE_BYTES = MAX_FILE_SIZE_MB * 1024 * 1024

SUPPORTED_EXTENSIONS = {".pdf", ".docx"}


def _get_file_bytes(uploaded_file) -> bytes:
    """Read uploaded file contents without permanently storing them."""
    try:
        if hasattr(uploaded_file, "getvalue"):
            data = uploaded_file.getvalue()
        else:
            uploaded_file.seek(0)
            data = uploaded_file.read()
            uploaded_file.seek(0)

        if not isinstance(data, bytes):
            raise ValueError("The uploaded file could not be read.")

        return data

    except ValueError:
        raise
    except Exception as exc:
        raise ValueError(
            "Unable to read the uploaded file. Please upload it again."
        ) from exc


def validate_file(uploaded_file):
    """
    Validate the uploaded resume's presence, extension, and size.

    Returns:
        tuple[bool, str]: Validation status and explanatory message.
    """
    if uploaded_file is None:
        return False, "No file uploaded. Please select a resume."

    filename = getattr(uploaded_file, "name", "")

    if not filename or not str(filename).strip():
        return False, "The uploaded file has no valid filename."

    extension = Path(str(filename)).suffix.lower()

    if extension not in SUPPORTED_EXTENSIONS:
        return False, "Unsupported file type. Please upload a PDF or DOCX."

    try:
        # Streamlit UploadedFile supports getvalue(), which also
        # avoids changing the current file position.
        if hasattr(uploaded_file, "getvalue"):
            file_size = len(uploaded_file.getvalue())
        else:
            uploaded_file.seek(0, 2)
            file_size = uploaded_file.tell()
            uploaded_file.seek(0)

    except Exception:
        return False, "Unable to determine file size. Please upload again."

    if file_size == 0:
        return False, "The uploaded file is empty."

    if file_size > MAX_FILE_SIZE_BYTES:
        return False, (
            f"File size ({file_size / (1024 * 1024):.2f} MB) "
            f"exceeds the {MAX_FILE_SIZE_MB} MB limit."
        )

    return True, "File validation successful."


def extract_text_from_pdf(file_stream) -> str:
    """Extract selectable text from every page of a PDF resume."""
    try:
        if isinstance(file_stream, bytes):
            pdf_data = file_stream
        else:
            pdf_data = _get_file_bytes(file_stream)

        reader = PdfReader(BytesIO(pdf_data), strict=False)

        if reader.is_encrypted:
            try:
                if reader.decrypt(""):
                    pass
                else:
                    raise ValueError(
                        "The PDF is password-protected. "
                        "Please upload an unlocked PDF."
                    )
            except ValueError:
                raise
            except Exception as exc:
                raise ValueError(
                    "Unable to open the encrypted PDF."
                ) from exc

        extracted_pages = []

        for page_number, page in enumerate(reader.pages, start=1):
            try:
                page_text = page.extract_text()
            except Exception as exc:
                raise ValueError(
                    f"Could not extract text from PDF page {page_number}."
                ) from exc

            if page_text and page_text.strip():
                extracted_pages.append(page_text.strip())

        full_text = "\n\n".join(extracted_pages).strip()

        if not full_text:
            raise ValueError(
                "No selectable text was found in this PDF. "
                "It may be scanned or image-based. "
                "Please upload a text-based PDF or use OCR."
            )

        return full_text

    except ValueError:
        raise
    except Exception as exc:
        raise ValueError(
            "Unable to parse the PDF. Please check that it is valid "
            "and not corrupted."
        ) from exc


def extract_text_from_docx(file_stream) -> str:
    """Extract text from DOCX paragraphs, tables, headers, and footers."""
    try:
        if isinstance(file_stream, bytes):
            docx_data = file_stream
        else:
            docx_data = _get_file_bytes(file_stream)

        document = Document(BytesIO(docx_data))
        text_chunks = []

        # Extract paragraphs.
        for paragraph in document.paragraphs:
            text = paragraph.text.strip()
            if text:
                text_chunks.append(text)

        # Extract tables, including resumes formatted as tables.
        for table in document.tables:
            for row in table.rows:
                cell_texts = [
                    cell.text.strip()
                    for cell in row.cells
                    if cell.text.strip()
                ]

                if cell_texts:
                    text_chunks.append(" | ".join(cell_texts))

        # Extract section headers and footers.
        for section in document.sections:
            for part in (section.header, section.footer):
                for paragraph in part.paragraphs:
                    text = paragraph.text.strip()
                    if text:
                        text_chunks.append(text)

                for table in part.tables:
                    for row in table.rows:
                        cell_texts = [
                            cell.text.strip()
                            for cell in row.cells
                            if cell.text.strip()
                        ]
                        if cell_texts:
                            text_chunks.append(" | ".join(cell_texts))

        full_text = "\n".join(text_chunks).strip()

        if not full_text:
            raise ValueError(
                "No readable text was found in the DOCX file."
            )

        return full_text

    except ValueError:
        raise
    except Exception as exc:
        raise ValueError(
            "Unable to parse the DOCX file. Please check that it is "
            "a valid, uncorrupted DOCX document."
        ) from exc


def parse_resume(uploaded_file) -> str:
    """
    Validate and extract raw text from an uploaded PDF or DOCX resume.

    Returns:
        str: Extracted resume text, ready for cleaning and analysis.

    Raises:
        ValueError: If validation or text extraction fails.
    """
    is_valid, message = validate_file(uploaded_file)

    if not is_valid:
        raise ValueError(message)

    filename = str(uploaded_file.name)
    extension = Path(filename).suffix.lower()

    # Read the file once and keep processing in memory.
    file_data = _get_file_bytes(uploaded_file)

    if extension == ".pdf":
        return extract_text_from_pdf(file_data)

    if extension == ".docx":
        return extract_text_from_docx(file_data)

    raise ValueError("Unsupported file format. Upload a PDF or DOCX.")
